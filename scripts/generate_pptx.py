import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette: Deep Tech Blue, Sky Blue, Warm Orange/Gold, Clean White/Gray
    DARK_BLUE = RGBColor(15, 32, 67)      # #0F2043
    MID_BLUE = RGBColor(31, 78, 121)      # #1F4E79
    ACCENT_CYAN = RGBColor(0, 168, 204)   # #00A8CC
    LIGHT_BG = RGBColor(245, 247, 250)    # #F5F7FA
    WHITE = RGBColor(255, 255, 255)
    CARD_BG = RGBColor(255, 255, 255)
    TEXT_MAIN = RGBColor(30, 41, 59)
    TEXT_MUTED = RGBColor(100, 116, 139)
    GOLD = RGBColor(217, 119, 6)

    def add_header(slide, title_text, category_text="第八届湖南省大学生智能导航科技创新大赛 · 科技创意类"):
        # Header banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT_CYAN
        p0.font.name = "SimHei"

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = DARK_BLUE
        p1.font.name = "SimHei"
        p1.space_before = Pt(4)

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=None):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    # ==================== SLIDE 1: COVER ====================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_BLUE
    bg1.line.fill.background()

    tb = slide1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(11.0), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "第八届湖南省大学生智能导航科技创新大赛 · 科技创意类作品"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.font.name = "SimHei"

    p = tf.add_paragraph()
    p.text = "银发导航智能体"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "SimHei"
    p.space_before = Pt(16)

    p = tf.add_paragraph()
    p.text = "基于多Agent协同的老年人安心出行伴侣"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(224, 231, 255)
    p.font.name = "SimHei"
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = "深度融合北斗卫星高精时空感知 · 适老化微地形代价路由 · 亲情双向守护闭环"
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(148, 163, 184)
    p.font.name = "SimSun"
    p.space_before = Pt(24)

    p = tf.add_paragraph()
    p.text = "主办单位：湖南省教育厅  |  承办单位：国防科技大学、测控与导航技术联合工程研究中心、湘南学院"
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(100, 116, 139)
    p.font.name = "SimSun"
    p.space_before = Pt(40)

    # ==================== SLIDE 2: PAIN POINTS & OPPORTUNITY ====================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "一、 立项背景与银发群体出行三大核心痛点")

    # 3 Cards
    cards_data = [
        ("痛点 1：传统导航算法不适老", "主流地图追求'距离最短/时间最短'，常推荐陡坡、长台阶、无盲道及复杂过街天桥，导致行动不便的老人寸步难行或面临跌倒风险。"),
        ("痛点 2：数字交互门槛高与弱智反问", "多层级菜单、小字号密集信息及多轮询问让老人望而却步；语音识别缺乏方言鲁棒性与语义宽容度，无法自主形成行动闭环。"),
        ("痛点 3：出行意外频发与子女监护悬空", "老人走失、偏航迷路、遭遇涉诈出行推销频发；子女远在异地无法感知实时轨迹，缺乏可靠的'北斗高精电子围栏'与双向安心响应机制。")
    ]
    for idx, (title, desc) in enumerate(cards_data):
        left = Inches(0.8 + idx * 3.95)
        add_card(slide2, left, Inches(1.8), Inches(3.75), Inches(4.8), CARD_BG, RGBColor(226, 232, 240))
        tb = slide2.shapes.add_textbox(left + Inches(0.25), Inches(2.0), Inches(3.25), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"0{idx+1}"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        p.font.name = "Arial"
        
        p = tf.add_paragraph()
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = DARK_BLUE
        p.font.name = "SimHei"
        p.space_before = Pt(8)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "SimSun"
        p.space_before = Pt(14)
        p.line_spacing = 1.3

    # ==================== SLIDE 3: SYSTEM ARCHITECTURE ====================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "二、 整体架构设计：北斗时空赋能的多Agent协同网络")

    add_card(slide3, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.0), CARD_BG, RGBColor(226, 232, 240))
    tb = slide3.shapes.add_textbox(Inches(1.05), Inches(1.9), Inches(5.2), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "5 大协同智能体家族分工"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE
    p.font.name = "SimHei"

    agents = [
        ("总调度智能体 (Main Agent)", "负责全域上下文感知、意图解析、任务并行分发与闭环决策。"),
        ("健康体能智能体 (Health Agent)", "读取慢病病史、关节炎/体能档案，输出最大步速与地形约束。"),
        ("北斗导航智能体 (BDS Nav Agent)", "对接北斗时空服务，计算微地形代价路由与高精走廊。"),
        ("气象感知智能体 (Weather Agent)", "监测高温暴雨结冰突变，动态触发路线修正或打车接驳。"),
        ("安全守护智能体 (Guardian Agent)", "负责电子围栏实时越界计算、异常滞留告警与反诈拦截。")
    ]
    for name, duty in agents:
        p = tf.add_paragraph()
        p.text = f"• {name}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = MID_BLUE
        p.font.name = "SimHei"
        p.space_before = Pt(8)
        
        p = tf.add_paragraph()
        p.text = f"   {duty}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "SimSun"

    add_card(slide3, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.0), CARD_BG, RGBColor(226, 232, 240))
    tb2 = slide3.shapes.add_textbox(Inches(7.05), Inches(1.9), Inches(5.2), Inches(4.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "端-边-云与事件总线底座"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE
    p.font.name = "SimHei"

    features = [
        ("微内核全序事件总线", "基于 Durable Event 架构，所有交互与决策落库可溯源、零死锁。"),
        ("单轮自主闭环机制", "长辈一句话发起，系统各 Agent 自动化协商交付，拒绝反问推诿。"),
        ("CGCS2000 国家大地坐标系解算", "严谨执行大地几何椭球投影与 NMEA-0183 卫星报文校验。"),
        ("双端透明协同机制", "长辈端极简操作 + 子女端全维数字孪生大屏，形成家庭守护纽带。")
    ]
    for title, desc in features:
        p = tf2.add_paragraph()
        p.text = f"✓ {title}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        p.font.name = "SimHei"
        p.space_before = Pt(10)
        
        p = tf2.add_paragraph()
        p.text = f"   {desc}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "SimSun"

    # ==================== SLIDE 4: BEIDOU INTEGRATION & ALGORITHMS ====================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "三、 北斗深度融合：高精时空基准与适老微地形代价模型")

    bds_items = [
        ("亚米级高精时空接入", "深度接入中国北斗地基增强网与卫星授时，达到亚米级车道与步道级精准定位，有效抵御城市峡谷遮挡与多路径效应。"),
        ("微地形通行代价方程 (Cost Formula)", "构建适老代价模型：Cost = L * (1 + W_slope * Slope^2 + W_step * N_steps - W_shade * Shade)；对台阶施加 10 倍惩罚权重，彻底杜绝长坡与陡梯。"),
        ("动态安全走廊 (BDS Corridor)", "依托北斗亚米级轨迹建立 25 米半径球面动态廊道缓冲带；一旦偏离即刻启动声控指引，杜绝老人迷失。"),
        ("北斗特种通信与应急概念联动", "预置北斗短报文与一键 SOS 救援信标，在公网信号盲区仍可发射带高精经纬度的求助脉冲。")
    ]
    for idx, (title, desc) in enumerate(bds_items):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.8 + row * 2.55)
        add_card(slide4, left, top, Inches(5.7), Inches(2.35), CARD_BG, RGBColor(226, 232, 240))
        tb = slide4.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), Inches(5.2), Inches(1.95))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = DARK_BLUE
        p.font.name = "SimHei"
        
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "SimSun"
        p.space_before = Pt(8)
        p.line_spacing = 1.3

    # ==================== SLIDE 5: ELDER NAVIGATION UI ====================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "四、 适老化极简交互：实景地标导航与大白话语音直达")

    ui_cards = [
        ("大白话一键语音直达", "无需手动输入起点与参数，长按大麦克风说出'想去烈士公园走走，腿有点酸'，智能体自动提取意图与身体状态。"),
        ("地标式路口实景引导卡片", "摒弃传统'200米后向东北方向转弯'等抽象术语，采用'过益丰大药房右转'、'顺着林荫平道直走'等醒目地标提示。"),
        ("北斗卫星遥测状态感知", "界面常驻 BDS 卫星遥测条（如 12 颗 BDS 亚米级锁定），给予老人和家属极强的高科技确定感与安全感。"),
        ("声控偏航与迷路自动纠偏", "监测到步道偏离时，不发出刺耳蜂鸣，而是通过温柔和蔼的长辈语调进行大白话原路纠偏。")
    ]
    for idx, (title, desc) in enumerate(ui_cards):
        left = Inches(0.8 + idx * 2.95)
        add_card(slide5, left, Inches(1.8), Inches(2.8), Inches(4.8), CARD_BG, RGBColor(226, 232, 240))
        tb = slide5.shapes.add_textbox(left + Inches(0.18), Inches(2.0), Inches(2.44), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = DARK_BLUE
        p.font.name = "SimHei"
        
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "SimSun"
        p.space_before = Pt(12)
        p.line_spacing = 1.3

    # ==================== SLIDE 6: GUARDIAN HUB ====================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "五、 子女端安心守护大屏：北斗电子围栏与亲情双向闭环")

    gh_items = [
        ("实时数字孪生轨迹看板", "子女端全屏呈现父母实时位置、行进航向、当前步速及剩余行程，全天行程自动绘制平滑时空轨迹。"),
        ("多级多边形北斗电子围栏", "支持绘制'常住小区-菜市场-社区医院'安全活动区；出圈毫秒级触发微信/App推送，越界风险无死角。"),
        ("异常静止与超时滞留检测", "在非长椅休闲区静止超 25 分钟时，系统自动判定疑似身体不适或摔倒，触发分级警报与子女代呼叫。"),
        ("一键报平安与代办托管", "老人到站一键触达'已平安抵达'语音卡片；子女可远程代订适老专车或预约就医挂号，形成孝心闭环。")
    ]
    for idx, (title, desc) in enumerate(gh_items):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.8 + row * 2.55)
        add_card(slide6, left, top, Inches(5.7), Inches(2.35), CARD_BG, RGBColor(226, 232, 240))
        tb = slide6.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), Inches(5.2), Inches(1.95))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = DARK_BLUE
        p.font.name = "SimHei"
        
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "SimSun"
        p.space_before = Pt(8)
        p.line_spacing = 1.3

    # ==================== SLIDE 7: EMERGENCY & GREEN CHANNEL ====================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "六、 突发风险主动防御：防诈拦截与就医急救绿通秒级重划")

    sc_cards = [
        ("涉诈与高危行程前置拦截", "对长辈前往偏远保健品推销窝点、异常高额消费场所进行主动风险警示与子女强提醒，从源头切断银发诈骗。"),
        ("突发急症就医绿通秒级重划", "途中长辈说'胸口发闷、喘不上气'，系统立即中断原路线，北斗导航毫秒级锁定就近三甲医院急诊通道并通报子女。"),
        ("跌倒与走失一键紧急呼救", "一键触发 SOS，高精度经纬度即时同步给子女、网格员与社区紧急联系人，建立立体应急救助网络。")
    ]
    for idx, (title, desc) in enumerate(sc_cards):
        left = Inches(0.8 + idx * 3.95)
        add_card(slide7, left, Inches(1.8), Inches(3.75), Inches(4.8), CARD_BG, RGBColor(226, 232, 240))
        tb = slide7.shapes.add_textbox(left + Inches(0.25), Inches(2.0), Inches(3.25), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = DARK_BLUE
        p.font.name = "SimHei"

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "SimSun"
        p.space_before = Pt(14)
        p.line_spacing = 1.3

    # ==================== SLIDE 8: HUNAN LOCALIZED DEMONSTRATION ====================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "七、 典型示范场景：湖南本地化特色实践")

    demo_cards = [
        ("示范场景一：湖南烈士公园休闲散步", "• 针对膝关节炎老人，规避烈士纪念塔陡长台阶，优选年嘉湖沿线无障碍平缓绿道。\n• 结合气象感知智能体，实时播报午后阵雨，自动引导至附近湖畔避雨凉亭。"),
        ("示范场景二：中南大学湘雅医院就医全流程", "• 智能体串联挂号号源、无障碍接驳与门诊导航。\n• 引导至湘雅路无障碍电梯通道，避开天桥施工拥堵，五页计划书一键交付子女确认。"),
        ("示范场景三：湖南省人民医院途中突发急症", "• 老人途中突发不适，系统秒级重划切换至省人民医院急诊绿通，北斗短报文脉冲广播至子女。"),
        ("示范场景四：橘子洲景区全景亲情守护", "• 设置橘子洲景区外轮廓安全围栏，精准感知老人观光小火车换乘，子女端随时随地安心查看。")
    ]
    for idx, (title, desc) in enumerate(demo_cards):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.8 + row * 2.55)
        add_card(slide8, left, top, Inches(5.7), Inches(2.35), CARD_BG, RGBColor(226, 232, 240))
        tb = slide8.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), Inches(5.2), Inches(1.95))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = DARK_BLUE
        p.font.name = "SimHei"
        
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "SimSun"
        p.space_before = Pt(8)
        p.line_spacing = 1.3

    # ==================== SLIDE 9: BUSINESS & SOCIAL IMPACT ====================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "八、 商业落地闭环与北斗规模化应用战略价值")

    biz_items = [
        ("B2C 孝心经济订阅模式", "提供基础适老出行免费，高级北斗安全走廊与突发急救秒级联动收取家庭安心年费（约 199元/年），打造高粘性家庭守护生态。"),
        ("B2B 康养硬件联合预装", "携手智能拐杖、老人手环及银发机顶盒厂商，提供北斗适老微地形导航 SDK 与守护后台，实现软硬一体赋能。"),
        ("B2G 智慧民政公共采购", "对接湖南省民政与老龄委'银发助老工程'，作为社区网格化养老与孤寡老人安全防走失的标准数字化基础设施。"),
        ("国家战略与产业价值", "契合《积极应对人口老龄化国家战略》，有力推动北斗高精度大众消费级应用，助力湖南'北斗+'产业规模化提质跃升。")
    ]
    for idx, (title, desc) in enumerate(biz_items):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.8 + row * 2.55)
        add_card(slide9, left, top, Inches(5.7), Inches(2.35), CARD_BG, RGBColor(226, 232, 240))
        tb = slide9.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), Inches(5.2), Inches(1.95))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = DARK_BLUE
        p.font.name = "SimHei"
        
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "SimSun"
        p.space_before = Pt(8)
        p.line_spacing = 1.3

    # ==================== SLIDE 10: CONCLUSION ====================
    slide10 = prs.slides.add_slide(blank_layout)
    bg10 = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg10.fill.solid()
    bg10.fill.fore_color.rgb = DARK_BLUE
    bg10.line.fill.background()

    tb = slide10.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "让中国北斗之光，照亮每一位长辈的安心归家路"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "SimHei"

    p = tf.add_paragraph()
    p.text = "北斗高精时空 · 多Agent协同闭环 · 适老化科技向善"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.font.name = "SimHei"
    p.space_before = Pt(18)

    p = tf.add_paragraph()
    p.text = "第八届湖南省大学生智能导航科技创新大赛 · 科技创意类参赛作品汇报"
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(148, 163, 184)
    p.font.name = "SimSun"
    p.space_before = Pt(30)

    prs.save(output_path)
    print(f"Generated PPTX: {output_path} ({os.path.getsize(output_path)} bytes)")

if __name__ == '__main__':
    deck_path = r'c:\Users\lenovo\Desktop\develop\laoyouji\docs\competition\银发导航智能体：基于多Agent协同的老年人安心出行伴侣_作品介绍PPT.pptx'
    desktop_deck_path = r'c:\Users\lenovo\Desktop\银发导航智能体：基于多Agent协同的老年人安心出行伴侣_作品介绍PPT.pptx'
    create_deck(deck_path)
    shutil.copy2(deck_path, desktop_deck_path)
    print("Copied PPTX to Desktop successfully!")
