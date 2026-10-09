# -*- coding: utf-8 -*-
"""
scripts/patch_slide17_pro.py
专业化重构 Slide 17：
1. 解决标题框宽度过大遮挡右侧标签的问题（width从10728655缩小到7500000）；
2. 标题文字精炼大气，彻底避免重叠；
3. 左侧卡片文本彻底消除口语化，转换为高规格的工科/学术级专业术语体系。
"""
import os
import pptx

def patch_slide17(prs_path):
    prs = pptx.Presentation(prs_path)
    s17 = prs.slides[16] # Slide 17
    
    # 1. 标题形状 Shape 1
    s1 = s17.shapes[1]
    s1.width = 7500000 # 限制宽度，防止与右侧 Shape 5/6 (left=8470900) 重叠
    s1.text_frame.paragraphs[0].text = "方案特色与展望：纯软件架构的高可用普惠范式"
    
    # 2. 左侧卡片 Shape 3
    s3 = s17.shapes[3]
    tf = s3.text_frame
    
    # P0: 标题
    tf.paragraphs[0].text = "【方案核心架构特色与普惠价值】"
    
    # P1: 分点1 标题
    tf.paragraphs[1].text = "①【纯软件智能体形态与零边际部署成本】"
    # P2: 分点1 内容
    tf.paragraphs[2].text = "立足科技创意类纯软件智能体形态，完全解耦专用外设硬件依赖，依托商用智能终端原生传感器实现零边际硬件成本的普惠智慧助老。"
    
    # P3: 分点2 标题
    tf.paragraphs[3].text = "②【终端多源传感器自适应感知融合】"
    # P4: 分点2 内容
    tf.paragraphs[4].text = "深度挖掘商用智能手机内置北斗多频基带、MEMS气压计与六轴IMU惯导时序数据，依靠适老抗微动滤波与自适应约束算法达成分米级高精解算。"
    
    # P5: 分点3 标题
    tf.paragraphs[5].text = "③【端云异构算力解耦与轻量化调度】"
    # P6: 分点3 内容
    tf.paragraphs[6].text = "高并发多Agent全序协商与微地形代价图求解由云端微内核引擎承载，终端仅负责轻量化SVG/Canvas渲染与心智流展示，显著降低算力与功耗开销。"
    
    # P7: 分点4 标题
    tf.paragraphs[7].text = "④【微空间无障碍数字标准与生态外延】"
    # P8: 分点4 内容
    tf.paragraphs[8].text = "提炼人行尺度微拓扑高程建图规范与适老代价路由模型，未来可标准化输出赋能城市人行无障碍微步道规划与社区养老数字化服务底座。"
    
    # P9: 脚注总结
    tf.paragraphs[9].text = "📌 核心范式：纯软件智能体形态 · 零硬件部署边际成本 · 微空间高精算法赋能"
    
    prs.save(prs_path)
    print(f"Slide 17 successfully patched: {prs_path}")

if __name__ == "__main__":
    targets = [
        r"C:\Users\lenovo\Desktop\科技创意类+银发导航智能体：基于多Agent协同的老年人安心出行伴侣+杨俊\银发导航智能体：基于多Agent协同的老年人安心出行伴侣+ppt.pptx",
        r"C:\Users\lenovo\Desktop\银发导航智能体_科幻风格背景版(1).pptx"
    ]
    for t in targets:
        patch_slide17(t)
