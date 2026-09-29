import re
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def md_to_docx(md_path, docx_path):
    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base normal style font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'SimSun'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(51, 51, 51)
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        rows_data = []
        for l in tbl_lines:
            parts = [p.strip() for p in l.strip().strip('|').split('|')]
            # Skip separator rows like |---|---|
            if all(re.match(r'^[-: ]+$', p) for p in parts if p):
                continue
            rows_data.append(parts)
        if not rows_data:
            return
        num_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for r_idx, row in enumerate(rows_data):
            for c_idx, cell_text in enumerate(row):
                if c_idx < num_cols:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = cell_text
                    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                    for p in cell.paragraphs:
                        p.paragraph_format.line_spacing = 1.15
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.space_before = Pt(2)
                        for r in p.runs:
                            r.font.name = 'SimSun'
                            r.font.size = Pt(9.5)
                            r._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
                            if r_idx == 0:
                                r.font.bold = True
                                r.font.color.rgb = RGBColor(255, 255, 255)
                    if r_idx == 0:
                        set_cell_background(cell, '1F4E79') # Navy blue
                    elif r_idx % 2 == 1:
                        set_cell_background(cell, 'F2F4F7') # Alternating light gray
        doc.add_paragraph() # space after table

    def flush_code(c_lines):
        if not c_lines:
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.line_spacing = 1.15
        
        # Put into a 1x1 styled table to simulate a code block card
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, 'F0F4F8')
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        cell.text = "\n".join(c_lines)
        for par in cell.paragraphs:
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.15
            for r in par.runs:
                r.font.name = 'Consolas'
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(40, 60, 80)
                r._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
        doc.add_paragraph()

    for line in lines:
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                flush_code(code_lines)
                code_lines = []
                in_code_block = False
            else:
                if in_table:
                    flush_table(table_lines)
                    table_lines = []
                    in_table = False
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line.rstrip('\r\n'))
            continue

        # Handle tables
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            flush_table(table_lines)
            table_lines = []
            in_table = False

        if not stripped:
            continue

        # Horizontal rule
        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            continue

        # Headings
        if stripped.startswith('# '):
            text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(10)
            run = p.add_run(text)
            run.font.name = 'SimHei'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(31, 78, 121)
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        elif stripped.startswith('## '):
            text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.font.name = 'SimHei'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(31, 78, 121)
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        elif stripped.startswith('### '):
            text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.name = 'SimHei'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(46, 117, 182)
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        elif stripped.startswith('#### '):
            text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.name = 'SimHei'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(50, 50, 50)
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

        # List items
        elif stripped.startswith('- ') or stripped.startswith('* '):
            text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.25
            add_formatted_text(p, text)

        elif re.match(r'^\d+\.\s', stripped):
            match = re.match(r'^\d+\.\s', stripped)
            text = stripped[len(match.group(0)):].strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.25
            add_formatted_text(p, text)

        # Blockquote
        elif stripped.startswith('> '):
            text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.2
            run = p.add_run("“ " + text + " ”")
            run.font.italic = True
            run.font.color.rgb = RGBColor(80, 80, 80)

        # Standard paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.first_line_indent = Inches(0.28) # 2 characters indent
            add_formatted_text(p, stripped)

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code(code_lines)

    doc.save(docx_path)
    print(f"Generated: {docx_path} ({os.path.getsize(docx_path)} bytes)")

def add_formatted_text(paragraph, text):
    # Match **bold** or regular text
    tokens = re.split(r'(\*\*.*?\*\*)', text)
    for token in tokens:
        if token.startswith('**') and token.endswith('**') and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.font.bold = True
            run.font.name = 'SimSun'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
        else:
            if token:
                run = paragraph.add_run(token)
                run.font.name = 'SimSun'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

if __name__ == '__main__':
    md_file = r'c:\Users\lenovo\Desktop\develop\laoyouji\docs\competition\01_PROJECT_PROPOSAL_REPORT.md'
    docx_file = r'c:\Users\lenovo\Desktop\develop\laoyouji\docs\competition\银发导航智能体：基于多Agent协同的老年人安心出行伴侣.docx'
    root_docx_file = r'c:\Users\lenovo\Desktop\develop\laoyouji\银发导航智能体：基于多Agent协同的老年人安心出行伴侣.docx'
    
    md_to_docx(md_file, docx_file)
    import shutil
    shutil.copy2(docx_file, root_docx_file)
    print("Done generating Word documents!")
